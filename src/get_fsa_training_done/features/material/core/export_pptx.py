"""Write a deck into the FSA template, slot by slot.

Each slide is added from the template's own layout, so the background, the
copyright line, the slide number and the fonts are the template's rather than an
imitation; each slot's content goes into the placeholder `layouts.py` maps it to.
Speaker notes, and the answers a check slide folds away, become the slide's notes.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from . import assets, i18n, rich
from . import deck as deck_mod
from . import layouts as L
from .dom import Document

BASE = Path(__file__).resolve().parents[1] / "theme" / "fsa_base.pptx"
MONO = "Consolas"
SANS = "Arial"
CODE_INK = "1F2328"
WRONG = "C62828"
RIGHT = "2E7D32"
LINE_HEIGHT = 1.25


@dataclass
class Outcome:
    slides: int = 0
    #: (slide label, message): what could not be carried over faithfully.
    warnings: list[tuple[str, str]] = field(default_factory=list)


def export(document: Document, out: Path, *, base: Path = BASE) -> Outcome:
    from pptx import Presentation

    prs = Presentation(str(base))
    by_name = {layout.name: layout for layout in prs.slide_layouts}
    slides = deck_mod.read(document)
    titles = deck_mod.section_titles(slides)
    outcome = Outcome()
    section = 0

    for slide in slides:
        if slide.layout is None:
            outcome.warnings.append((slide.where(), "unknown layout; slide skipped"))
            continue
        layout = by_name[slide.layout.pptx_name]
        target = prs.slides.add_slide(layout)
        placeholders = {shape.placeholder_format.idx: shape for shape in target.placeholders}
        _inherit_shapes(placeholders, layout)
        if slide.role_kind == L.SECTION:
            section += 1

        for spec in slide.layout.slots:
            element = slide.slots.get(spec.name)
            shape = placeholders.get(spec.idx) if spec.idx is not None else None
            if shape is not None:
                _place(shape, spec.box, prs)
            if element is None:
                if spec.default and shape is not None:
                    words = i18n.t(document.lang, spec.default)
                    default = [rich.Block("para", runs=[rich.Run(words)])]
                    _write_text(shape, default, spec, scale=L.fit_scale(spec, words))
                elif shape is not None:
                    _remove(shape)
                continue
            content = rich.blocks(element)
            if spec.name == "title" and slide.role_kind == L.SECTION:
                content = [
                    rich.Block(
                        "para",
                        runs=[rich.Run(f"{section:02d}.\n" + element.text(keep_breaks=True))],
                    )
                ]
            _fill(target, shape, spec, content, slide, document, outcome)

        if slide.role_kind == L.TOC:
            _toc(target, placeholders, titles, outcome, slide)

        notes = rich.notes_paragraphs(slide.notes)
        for answer in rich.answers(slide.element):
            notes.append(f"Answer: {answer}")
        if notes:
            target.notes_slide.notes_text_frame.text = "\n".join(notes)
        outcome.slides += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return outcome


# --- Placing --------------------------------------------------------------


def _emu(box: L.Box, prs) -> tuple[int, int, int, int]:
    width, height = prs.slide_width, prs.slide_height
    return (
        round(box.left / 100 * width),
        round(box.top / 100 * height),
        round(box.width / 100 * width),
        round(box.height / 100 * height),
    )


def _place(shape, box: L.Box, prs) -> None:
    shape.left, shape.top, shape.width, shape.height = _emu(box, prs)


def _remove(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


#: The parts of a placeholder's shape a slide inherits from its layout. PowerPoint
#: applies the inheritance itself; other viewers (Keynote, Quick Look, Google
#: Slides) often do not, so each is copied onto the slide explicitly.
INHERITED = ("a:prstGeom", "a:custGeom", "a:noFill", "a:solidFill", "a:gradFill", "a:ln")


def _inherit_shapes(placeholders: dict, layout) -> None:
    import copy

    from pptx.oxml.ns import qn

    sources = {shape.placeholder_format.idx: shape for shape in layout.placeholders}
    for idx, shape in placeholders.items():
        source = sources.get(idx)
        if source is None:
            continue
        src_pr = source._element.spPr
        dst_pr = shape._element.spPr
        if dst_pr.find(qn("a:xfrm")) is None and src_pr.find(qn("a:xfrm")) is not None:
            dst_pr.append(copy.deepcopy(src_pr.find(qn("a:xfrm"))))
        for tag in INHERITED:
            element = src_pr.find(qn(tag))
            if element is not None and dst_pr.find(qn(tag)) is None:
                dst_pr.append(copy.deepcopy(element))


def text_of(content) -> str:
    """The plain text a fit estimate measures."""
    return "\n".join(block.text().replace("\v", "\n") for block in content if block.runs)


def _toc(target, placeholders: dict, titles: list[str], outcome: Outcome, slide) -> None:
    """Each row: the template's numbered circle and the section title beside it.

    The circle is drawn as a shape of its own, in the template's geometry, fill
    and ring: viewers other than PowerPoint do not paint a placeholder's
    inherited fill, and a number with no circle reads as a stray digit.
    """
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Pt

    for row, ((number_idx, number_box), (title_idx, title_box)) in enumerate(L.TOC_ROWS):
        number, title = placeholders.get(number_idx), placeholders.get(title_idx)
        if row >= len(titles):
            for shape in (number, title):
                if shape is not None:
                    _remove(shape)
            continue
        if number is not None:
            left, top, width, height = number.left, number.top, number.width, number.height
            _remove(number)
            circle = target.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
            circle.fill.solid()
            circle.fill.fore_color.rgb = RGBColor.from_string(L.GREEN)
            circle.line.color.rgb = RGBColor.from_string(L.WHITE)
            circle.line.width = Pt(2.25)
            circle.shadow.inherit = False
            _write_text(
                circle,
                [rich.Block("para", runs=[rich.Run(str(row + 1))])],
                _toc_spec(number_box, "center"),
            )
        if title is not None:
            _write_text(
                title,
                [rich.Block("para", runs=[rich.Run(titles[row])])],
                _toc_spec(title_box, "left"),
            )
    if len(titles) > len(L.TOC_ROWS):
        outcome.warnings.append(
            (slide.where(), f"{len(titles)} sections; the template lists {len(L.TOC_ROWS)}")
        )


def _toc_spec(box: L.Box, align: str) -> L.Slot:
    return L.Slot(
        "toc",
        None,
        box,
        L.TEXT,
        size=L.TOC_SIZE,
        colour=L.WHITE,
        italic=True,
        align=align,
        anchor="middle",
    )


CAPTION_INK = "5B6370"
CAPTION_SCALE = 0.55


def _fill(target, shape, spec: L.Slot, content, slide, document: Document, outcome) -> None:
    """Text before the first graphic at the top, the graphics, then what follows them.

    The order is the HTML's: a figure's caption stays under its picture.
    """
    from pptx.util import Emu

    is_graphic = [block.kind in ("image", "table") for block in content]
    first = is_graphic.index(True) if any(is_graphic) else len(content)
    last = len(content) - 1 - is_graphic[::-1].index(True) if any(is_graphic) else len(content)
    before = [block for block in content[:first] if block.kind != "caption"]
    graphics = [block for block, graphic in zip(content, is_graphic, strict=True) if graphic]
    after = list(content[last + 1 :]) if graphics else []
    if not graphics:
        before = [block for block in content if block.kind != "caption"]
        after = [block for block in content if block.kind == "caption"]

    prs = target.part.package.presentation_part.presentation
    left = round(spec.box.left / 100 * prs.slide_width)
    top = round(spec.box.top / 100 * prs.slide_height)
    width = round(spec.box.width / 100 * prs.slide_width)
    height = round(spec.box.height / 100 * prs.slide_height)

    head = 0
    if before:
        if shape is None:
            shape = target.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
        _write_text(
            shape, before, spec, scale=L.fit_scale(spec, text_of(before)), lang=document.lang
        )
        head = min(height * 0.6, _text_height_pt(before, spec) * 12700) if graphics else 0
    elif shape is not None:
        _remove(shape)

    foot = 0
    if after:
        foot = min(height * 0.3, _text_height_pt(after, spec) * 12700)
        box = target.shapes.add_textbox(
            Emu(left), Emu(top + height - int(foot)), Emu(width), Emu(int(foot))
        )
        _write_text(box, after, spec, lang=document.lang)

    if not graphics:
        return
    region = (left, top + int(head), width, height - int(head) - int(foot))
    slot_height = region[3] // len(graphics)
    for index, block in enumerate(graphics):
        box = (region[0], region[1] + index * slot_height, region[2], slot_height)
        if block.kind == "image":
            _picture(target, block, box, document, slide, spec, outcome)
        else:
            _table(target, block, box, spec)


def _text_height_pt(content, spec: L.Slot) -> float:
    total = 0.0
    for block in content:
        if block.kind in ("code", "diagram"):
            total += len(block.lines) * _code_size(spec) * LINE_HEIGHT + 8
        elif block.kind == "caption":
            total += (spec.size or 18) * CAPTION_SCALE * LINE_HEIGHT * 1.4 + 8
        else:
            total += _size(spec, block.level) * LINE_HEIGHT * 1.15
    return total


def _size(spec: L.Slot, level: int, scale: float = 1.0) -> float:
    size = (spec.size or 18) * scale
    return round(size * (0.82**level) * 2) / 2


def _code_size(spec: L.Slot) -> float:
    return L.code_size(spec)


# --- Text -------------------------------------------------------------------


ANCHORS = {"top": "TOP", "middle": "MIDDLE", "bottom": "BOTTOM"}
ALIGNS = {"left": "LEFT", "center": "CENTER"}


def _write_text(shape, content, spec: L.Slot, *, scale: float = 1.0, lang: str = "en") -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = getattr(MSO_ANCHOR, ANCHORS[spec.anchor])
    for paragraph in list(frame.paragraphs)[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    first = True
    colour = RGBColor.from_string(spec.colour)
    align = getattr(PP_ALIGN, ALIGNS[spec.align])

    def paragraph():
        nonlocal first
        if first:
            first = False
            para = frame.paragraphs[0]
            for run in list(para.runs):
                run._r.getparent().remove(run._r)
            return para
        return frame.add_paragraph()

    for block in content:
        if block.kind in ("code", "diagram"):
            if block.label:
                para = paragraph()
                _bullet(para, None)
                run = para.add_run()
                run.text = (
                    f"✗ {i18n.t(lang, 'Wrong')}"
                    if block.label == "wrong"
                    else f"✓ {i18n.t(lang, 'Right')}"
                )
                run.font.name = SANS
                run.font.bold = True
                run.font.size = Pt(_code_size(spec))
                run.font.color.rgb = RGBColor.from_string(
                    WRONG if block.label == "wrong" else RIGHT
                )
            for line in block.lines or [""]:
                para = paragraph()
                _bullet(para, None)
                para.line_spacing = 1.0
                run = para.add_run()
                run.text = line.replace("\t", "    ") or " "
                run.font.name = MONO
                run.font.size = Pt(_code_size(spec))
                run.font.color.rgb = RGBColor.from_string(
                    CODE_INK if spec.colour == L.BLACK else spec.colour
                )
            continue
        if block.kind == "caption":
            para = paragraph()
            para.alignment = PP_ALIGN.CENTER
            _bullet(para, None)
            run = para.add_run()
            run.text = block.text()
            run.font.name = spec.font
            run.font.size = Pt(round((spec.size or 18) * CAPTION_SCALE))
            run.font.color.rgb = RGBColor.from_string(CAPTION_INK)
            continue
        para = paragraph()
        para.alignment = align
        para.level = min(block.level, 4)
        _bullet(para, block.kind if block.kind in ("bullet", "number") else None, block.level)
        for item in block.runs:
            pieces = item.text.replace("\v", "\n").split("\n")
            for index, piece in enumerate(pieces):
                if index:
                    para.add_line_break()
                if piece:
                    _run(para, piece, item, spec, block.level, colour, scale)


def _run(para, text: str, item, spec: L.Slot, level: int, colour, scale: float) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    run = para.add_run()
    run.text = text
    run.font.name = MONO if item.code else spec.font
    run.font.size = Pt(_size(spec, level, scale))
    run.font.bold = item.bold or spec.bold
    run.font.italic = item.italic or spec.italic
    run.font.color.rgb = colour
    if item.bold and spec.colour == L.BLACK:
        run.font.color.rgb = RGBColor.from_string("034EA1")
    if item.href:
        run.hyperlink.address = item.href


def _bullet(paragraph, kind: str | None, level: int = 0) -> None:
    """Set a paragraph's bullet explicitly, whatever the placeholder's default."""
    from pptx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buClr"):
        for old in p_pr.findall(qn(tag)):
            p_pr.remove(old)
    indent = 228600  # 0.25in
    if kind is None:
        p_pr.set("marL", str(level * indent))
        p_pr.set("indent", "0")
        p_pr.append(p_pr.makeelement(qn("a:buNone"), {}))
        return
    p_pr.set("marL", str((level + 1) * indent))
    p_pr.set("indent", str(-indent))
    colour = p_pr.makeelement(qn("a:buClr"), {})
    srgb = colour.makeelement(qn("a:srgbClr"), {"val": "50B748" if level == 0 else "034EA1"})
    colour.append(srgb)
    p_pr.append(colour)
    if kind == "number":
        p_pr.append(p_pr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
    else:
        p_pr.append(p_pr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        p_pr.append(p_pr.makeelement(qn("a:buChar"), {"char": "•"}))


# --- Graphics ---------------------------------------------------------------


def _picture(target, block, box, document: Document, slide, spec, outcome) -> None:
    from pptx.util import Emu

    path, problem = assets.embeddable(document.path, block.src)
    if path is None:
        outcome.warnings.append((slide.where(), f"{problem}; image skipped"))
        return
    from PIL import Image

    with Image.open(path) as image:
        ratio = image.width / image.height
    left, top, width, height = box
    if width / height > ratio:
        fitted_w, fitted_h = int(height * ratio), height
    else:
        fitted_w, fitted_h = width, int(width / ratio)
    picture = target.shapes.add_picture(
        str(path),
        Emu(left + (width - fitted_w) // 2),
        Emu(top + (height - fitted_h) // 2),
        Emu(fitted_w),
        Emu(fitted_h),
    )
    if block.alt:
        picture._element.nvPicPr.cNvPr.set("descr", block.alt)


def _table(target, block, box, spec) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    rows = block.rows
    if not rows:
        return
    columns = max(len(row) for row in rows)
    left, top, width, height = box
    row_height = min(height // len(rows), Pt(28))
    graphic = target.shapes.add_table(
        len(rows), columns, Emu(left), Emu(top), Emu(width), Emu(row_height * len(rows))
    )
    table = graphic.table
    size = Pt(max(10, min(16, round((spec.size or 18) * 0.7))))
    for r, row in enumerate(rows):
        for c in range(columns):
            cell = table.cell(r, c)
            cell.text = row[c] if c < len(row) else ""
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = size
                    if r == 0 and block.header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if r == 0 and block.header:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string("034EA1")


__all__ = ["BASE", "Outcome", "export"]
