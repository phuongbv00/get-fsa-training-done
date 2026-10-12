#!/usr/bin/env python3
"""Build the FSA theme the material feature ships, from the FSA slide template.

Run by hand, once per template version — the source `.pptx` is not in the
repository, so this is deliberately not a `gen_*` script and CI never runs it:

    python scripts/material/build_theme.py ~/Downloads/Template_FSA_Training\\ Materials_v1.0.pptx

It writes, under `src/get_fsa_training_done/features/material/theme/`:

- `fsa_base.pptx` — the template's master and the layouts `core/layouts.py`
  maps, with every sample slide removed and every image recompressed. `material
  export pptx` adds slides to a copy of it, so an exported deck *is* the
  template rather than an imitation of it.
- `img/*.jpg` — the same backgrounds, named by role, for the HTML theme, plus
  `band.jpg`, the content layout's header band, for long-form documents.

Committing the theme is a deliberate exception to "a customer's template is
never committed": the owner chose a self-contained skill over fetching it.
Everything else in `theme/` (`fsa.css`, `deck.js`, `layouts.css`) is written by
hand or generated from `core/layouts.py`, never by this script.
"""

from __future__ import annotations

import argparse
import io
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from get_fsa_training_done.features.material.core import layouts as L  # noqa: E402

THEME = ROOT / "src" / "get_fsa_training_done" / "features" / "material" / "theme"
WIDTH = 1600
QUALITY = 78
#: The content layout's header band, as a fraction of the slide's height.
BAND = 0.145


def _recompress(blob: bytes, *, width: int = WIDTH) -> bytes:
    from PIL import Image

    image = Image.open(io.BytesIO(blob)).convert("RGB")
    if image.width > width:
        image = image.resize((width, round(image.height * width / image.width)), Image.LANCZOS)
    out = io.BytesIO()
    image.save(out, "JPEG", quality=QUALITY, optimize=True, progressive=True)
    return out.getvalue()


def _strip(source: Path) -> bytes:
    """The template with no slides and only the mapped layouts."""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    prs = Presentation(str(source))
    slide_ids = prs.slides._sldIdLst
    for sld_id in list(slide_ids):
        prs.part.drop_rel(sld_id.get(qn("r:id")))
        slide_ids.remove(sld_id)

    wanted = {layout.pptx_name for layout in L.LAYOUTS}
    master = prs.slide_masters[0]
    id_list = master._element.find(qn("p:sldLayoutIdLst"))
    found = set()
    for entry in list(id_list):
        rid = entry.get(qn("r:id"))
        name = master.part.related_part(rid).slide_layout.name
        if name in wanted:
            found.add(name)
            continue
        master.part.drop_rel(rid)
        id_list.remove(entry)
    missing = wanted - found
    if missing:
        raise SystemExit(f"the template has no layout named: {', '.join(sorted(missing))}")

    # The template embeds Calibri and Segoe UI. They are Microsoft's to license,
    # not ours to redistribute, and every machine that opens a deck has them or
    # a metric-compatible substitute.
    fonts = prs.part._element.find(qn("p:embeddedFontLst"))
    if fonts is not None:
        for rid in {el.get(qn("r:id")) for el in fonts.iter() if el.get(qn("r:id")) is not None}:
            prs.part.drop_rel(rid)
        prs.part._element.remove(fonts)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _shrink_media(pptx: bytes) -> tuple[bytes, dict[str, bytes]]:
    """Recompress every image in the package; return it and the media by name."""
    media: dict[str, bytes] = {}
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(pptx)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.startswith("ppt/media/") and info.filename.lower().endswith(
                (".jpg", ".jpeg", ".png")
            ):
                data = _recompress(data)
                media[info.filename] = data
            dst.writestr(info, data)
    return out.getvalue(), media


def _backgrounds(pptx: bytes) -> dict[str, bytes]:
    """Each mapped layout's background image, keyed by its `background` role."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(pptx))
    by_name = {layout.name: layout for layout in prs.slide_layouts}
    images: dict[str, bytes] = {}
    for layout in L.LAYOUTS:
        part = by_name[layout.pptx_name].part
        blobs = [
            rel.target_part.blob
            for rel in part.rels.values()
            if rel.reltype.endswith("/image") and not rel.is_external
        ]
        if not blobs:
            raise SystemExit(f"layout {layout.pptx_name!r} has no image")
        images.setdefault(layout.background, blobs[0])
    return images


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", help="the FSA Training Materials .pptx")
    args = parser.parse_args()

    source = Path(args.template).expanduser()
    if not source.is_file():
        print(f"ERROR: not found: {source}", file=sys.stderr)
        return 2

    from PIL import Image

    stripped = _strip(source)
    base, _media = _shrink_media(stripped)
    (THEME / "img").mkdir(parents=True, exist_ok=True)
    (THEME / "fsa_base.pptx").write_bytes(base)

    for role, blob in _backgrounds(base).items():
        (THEME / "img" / f"{role}.jpg").write_bytes(blob)

    content = Image.open(io.BytesIO(_backgrounds(base)["content"]))
    band = content.crop((0, 0, content.width, round(content.height * BAND)))
    band.save(THEME / "img" / "band.jpg", "JPEG", quality=QUALITY, optimize=True)

    total = sum(path.stat().st_size for path in THEME.rglob("*") if path.is_file())
    print(f"wrote {THEME.relative_to(ROOT)} ({total / 1024 / 1024:.1f} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
